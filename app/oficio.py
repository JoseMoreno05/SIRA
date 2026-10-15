"""Oficio de propuesta en Word y PDF. Siempre sale marcado como borrador para revisión humana."""
from __future__ import annotations

import io
import os
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

INSTITUCION = os.environ.get("SIRA_INSTITUCION", "Institución de demostración")
AZUL = (0x12, 0x35, 0x5B)
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]


def _fecha() -> str:
    h = date.today()
    return f"{h.day} de {MESES[h.month - 1]} de {h.year}"


def _filas(prop: dict) -> list[list[str]]:
    datos = prop["contenido"]["datos"]
    filas = []
    for m in prop["contenido"]["materias"]:
        d = datos[m["clave"]]
        filas.append([m["clave"], d["nombre"], str(m["grupos"]), str(m["cupo_estimado"]),
                      str(d["urgencia"]), m["prioridad"].capitalize()])
    return filas


ENCABEZADOS = ["Clave", "Materia", "Grupos", "Cupo estimado", "En riesgo alto", "Prioridad"]


def _origen(prop: dict) -> str:
    if prop["origen"] == "gemma":
        return (f"Redacción asistida por Gemma 4 ({prop['modelo']}); cifras calculadas por SIRA y verificadas "
                "contra la base antes de emitir este borrador.")
    return "Redacción con plantilla de SIRA (modo sin IA); cifras calculadas por SIRA."


def docx(prop: dict) -> bytes:
    c = prop["contenido"]
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)

    enc = doc.add_paragraph()
    enc.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = enc.add_run(f"{INSTITUCION}\nCiudad del Carmen, Campeche, a {_fecha()}")
    r.font.size = Pt(10)
    marca = doc.add_paragraph()
    marca.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    mr = marca.add_run("BORRADOR PARA REVISIÓN · generado con el prototipo SIRA")
    mr.bold = True
    mr.font.size = Pt(9)
    mr.font.color.rgb = RGBColor(0xA1, 0x5C, 0x07)

    doc.add_heading(f"Propuesta de oferta: {c['intersemestral']}", level=1)
    doc.add_paragraph("A la Jefatura del programa educativo:")
    doc.add_paragraph(c["resumen"])

    tabla = doc.add_table(rows=1, cols=len(ENCABEZADOS))
    tabla.style = "Light Grid Accent 1"
    for i, h in enumerate(ENCABEZADOS):
        tabla.rows[0].cells[i].text = h
    for fila in _filas(prop):
        celdas = tabla.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = v

    doc.add_heading("Justificación por materia", level=2)
    for m in c["materias"]:
        p = doc.add_paragraph()
        p.add_run(f"{m['clave']} · {c['datos'][m['clave']]['nombre']}. ").bold = True
        p.add_run(m["justificacion"])

    if c.get("descartadas"):
        doc.add_heading("Materias no propuestas", level=2)
        for d in c["descartadas"]:
            doc.add_paragraph(f"{d['clave']}: {d['motivo']}", style="List Bullet")

    cob = c["cobertura"]
    doc.add_heading("Alcance de los datos", level=2)
    doc.add_paragraph(
        f"Datos de {cob['tutores_sincronizados']} de {cob['tutores_total']} tutores y {cob['alumnos']} alumnos "
        f"sincronizados. {_origen(prop)}"
    )
    doc.add_paragraph("\n\n______________________________\nVo. Bo. Jefatura del programa educativo")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def pdf(prop: dict) -> bytes:
    c = prop["contenido"]
    buf = io.BytesIO()
    plantilla = SimpleDocTemplate(
        buf, pagesize=LETTER, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
        title=f"Propuesta de oferta: {c['intersemestral']}", author="SIRA", subject="Borrador de propuesta intersemestral",
        lang="es-MX",
    )
    base = getSampleStyleSheet()
    azul = colors.Color(*(v / 255 for v in AZUL))
    normal = ParagraphStyle("n", parent=base["Normal"], fontName="Helvetica", fontSize=10.5, leading=15, alignment=TA_LEFT)
    pequena = ParagraphStyle("p", parent=normal, fontSize=9, leading=12, textColor=colors.HexColor("#4A5565"))
    titulo = ParagraphStyle("t", parent=normal, fontName="Times-Bold", fontSize=17, leading=21, textColor=azul, spaceAfter=6)
    sub = ParagraphStyle("s", parent=normal, fontName="Times-Bold", fontSize=12.5, leading=16, textColor=azul, spaceBefore=10)
    marca = ParagraphStyle("m", parent=pequena, textColor=colors.HexColor("#A15C07"), fontName="Helvetica-Bold")

    h = [
        Paragraph(f"{INSTITUCION}<br/>Ciudad del Carmen, Campeche, a {_fecha()}", pequena),
        Paragraph("BORRADOR PARA REVISIÓN · generado con el prototipo SIRA", marca),
        Spacer(1, 10),
        Paragraph(f"Propuesta de oferta: {c['intersemestral']}", titulo),
        Paragraph("A la Jefatura del programa educativo:", normal), Spacer(1, 4),
        Paragraph(c["resumen"], normal), Spacer(1, 10),
    ]
    celdas = [[Paragraph(f"<b>{x}</b>", pequena) for x in ENCABEZADOS]] + [
        [Paragraph(v, pequena) for v in f] for f in _filas(prop)]
    t = Table(celdas, colWidths=[1.7 * cm, 6.0 * cm, 1.6 * cm, 2.3 * cm, 2.3 * cm, 2.0 * cm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6ECF3")),
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#C5CFDB")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    h += [t, Paragraph("Justificación por materia", sub)]
    for m in c["materias"]:
        h.append(Paragraph(f"<b>{m['clave']} · {c['datos'][m['clave']]['nombre']}.</b> {m['justificacion']}", normal))
        h.append(Spacer(1, 4))
    if c.get("descartadas"):
        h.append(Paragraph("Materias no propuestas", sub))
        for d in c["descartadas"]:
            h.append(Paragraph(f"{d['clave']}: {d['motivo']}", pequena))
    cob = c["cobertura"]
    h += [Paragraph("Alcance de los datos", sub),
          Paragraph(f"Datos de {cob['tutores_sincronizados']} de {cob['tutores_total']} tutores y {cob['alumnos']} "
                    f"alumnos sincronizados. {_origen(prop)}", pequena),
          Spacer(1, 40), Paragraph("______________________________<br/>Vo. Bo. Jefatura del programa educativo", normal)]
    plantilla.build(h)
    return buf.getvalue()


# ------------------------------------------------------------------ oficio semestral
ENCABEZADOS_SEM = ["Clave", "Materia", "Sem.", "Demanda esperada", "Recursan", "En riesgo alto", "Grupos"]


def _filas_sem(prop: dict) -> list[list[str]]:
    return [[m["clave"], m["nombre"], str(m["semestre"]), str(m["demanda_esperada"]), str(m["recursan"]),
             str(m["en_riesgo"]), str(m["grupos"]) if m["abre"] else "Sin grupo"]
            for m in prop["contenido"]["oferta"]["materias"]]


def _origen_sem(prop: dict) -> str:
    if prop["origen"] == "gemma":
        return (f"Resumen y observaciones redactados por Gemma 4 ({prop['modelo']}); grupos y cifras calculados "
                "por SIRA y verificados contra la base.")
    return "Redacción con plantilla de SIRA (modo sin IA); grupos y cifras calculados por SIRA."


def docx_semestral(prop: dict) -> bytes:
    c = prop["contenido"]
    o = c["oferta"]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    enc = doc.add_paragraph()
    enc.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    enc.add_run(f"{INSTITUCION}\nCiudad del Carmen, Campeche, a {_fecha()}").font.size = Pt(10)
    marca = doc.add_paragraph()
    marca.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    mr = marca.add_run("BORRADOR PARA REVISIÓN · generado con el prototipo SIRA")
    mr.bold = True
    mr.font.size = Pt(9)
    mr.font.color.rgb = RGBColor(0xA1, 0x5C, 0x07)
    doc.add_heading(f"Propuesta de oferta semestral: {o['nombre']}", level=1)
    doc.add_paragraph("A la Jefatura del programa educativo:")
    doc.add_paragraph(c["resumen"])
    if c["observaciones"]:
        doc.add_heading("Materias que requieren atención", level=2)
        nombres = {m["clave"]: m["nombre"] for m in o["materias"]}
        for ob in c["observaciones"]:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{nombres.get(ob['clave'], ob['clave'])}: ").bold = True
            p.add_run(ob["texto"])
    doc.add_heading("Oferta por materia", level=2)
    tabla = doc.add_table(rows=1, cols=len(ENCABEZADOS_SEM))
    tabla.style = "Light Grid Accent 1"
    for i, h in enumerate(ENCABEZADOS_SEM):
        tabla.rows[0].cells[i].text = h
    for fila in _filas_sem(prop):
        celdas = tabla.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = v
    cob = o["cobertura"]
    doc.add_heading("Alcance de los datos", level=2)
    doc.add_paragraph(
        f"Carga sugerida de {o['totales']['alumnos']} alumnos, con datos de {cob['tutores_sincronizados']} de "
        f"{cob['tutores_total']} tutores. No incluye alumnos de nuevo ingreso. Las materias condicionadas suponen "
        f"que el alumno acredita lo que cursa hoy y se ponderan por la tasa histórica. {_origen_sem(prop)}")
    doc.add_paragraph("\n\n______________________________\nVo. Bo. Jefatura del programa educativo")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def pdf_semestral(prop: dict) -> bytes:
    c = prop["contenido"]
    o = c["oferta"]
    buf = io.BytesIO()
    plantilla = SimpleDocTemplate(buf, pagesize=LETTER, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                                  bottomMargin=2 * cm, title=f"Propuesta de oferta semestral: {o['nombre']}",
                                  author="SIRA", subject="Borrador de propuesta semestral")
    base = getSampleStyleSheet()
    azul = colors.Color(*(v / 255 for v in AZUL))
    normal = ParagraphStyle("n", parent=base["Normal"], fontName="Helvetica", fontSize=10.5, leading=15)
    pequena = ParagraphStyle("p", parent=normal, fontSize=8.8, leading=11.5, textColor=colors.HexColor("#4A5565"))
    titulo = ParagraphStyle("t", parent=normal, fontName="Times-Bold", fontSize=17, leading=21, textColor=azul, spaceAfter=6)
    sub = ParagraphStyle("s", parent=normal, fontName="Times-Bold", fontSize=12.5, leading=16, textColor=azul, spaceBefore=10)
    marca = ParagraphStyle("m", parent=pequena, textColor=colors.HexColor("#A15C07"), fontName="Helvetica-Bold")
    nombres = {m["clave"]: m["nombre"] for m in o["materias"]}
    h = [Paragraph(f"{INSTITUCION}<br/>Ciudad del Carmen, Campeche, a {_fecha()}", pequena),
         Paragraph("BORRADOR PARA REVISIÓN · generado con el prototipo SIRA", marca), Spacer(1, 10),
         Paragraph(f"Propuesta de oferta semestral: {o['nombre']}", titulo),
         Paragraph("A la Jefatura del programa educativo:", normal), Spacer(1, 4), Paragraph(c["resumen"], normal)]
    if c["observaciones"]:
        h.append(Paragraph("Materias que requieren atención", sub))
        for ob in c["observaciones"]:
            h.append(Paragraph(f"<b>{nombres.get(ob['clave'], ob['clave'])}:</b> {ob['texto']}", normal))
            h.append(Spacer(1, 3))
    h.append(Paragraph("Oferta por materia", sub))
    celdas = [[Paragraph(f"<b>{x}</b>", pequena) for x in ENCABEZADOS_SEM]] + [
        [Paragraph(v, pequena) for v in f] for f in _filas_sem(prop)]
    t = Table(celdas, colWidths=[1.6 * cm, 5.6 * cm, 1.1 * cm, 2.3 * cm, 1.8 * cm, 2.1 * cm, 1.9 * cm], repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6ECF3")),
                           ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#C5CFDB")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    cob = o["cobertura"]
    h += [t, Paragraph("Alcance de los datos", sub),
          Paragraph(f"Carga sugerida de {o['totales']['alumnos']} alumnos, con datos de {cob['tutores_sincronizados']} "
                    f"de {cob['tutores_total']} tutores. No incluye alumnos de nuevo ingreso. Las materias condicionadas "
                    f"suponen que el alumno acredita lo que cursa hoy y se ponderan por la tasa histórica. "
                    f"{_origen_sem(prop)}", pequena),
          Spacer(1, 36), Paragraph("______________________________<br/>Vo. Bo. Jefatura del programa educativo", normal)]
    plantilla.build(h)
    return buf.getvalue()
